import pickle
import json
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from etl_processor import ETLProcessor

print("=== EXPORTANDO CACHÉ COMPLETO A FORMATOS GOOGLE DRIVE / SHEETS ===")

CACHE_FILE = "sql_cache.pkl"
processor = ETLProcessor()

if os.path.exists(CACHE_FILE):
    with open(CACHE_FILE, "rb") as f:
        cache_pack = pickle.load(f)

    summary_data = processor.get_summary(cache_pack, filters={"periodo": "ambos"})

    sedes_summary = summary_data.get("sedes_summary", [])
    pend_df = cache_pack.get("pendientes", pd.DataFrame())
    usuarios_df = cache_pack.get("usuarios", pd.DataFrame())

    # 1. Exportar a JSON legible para Google Drive
    json_path = "sql_cache_drive.json"
    
    # Formatear fechas a string para JSON
    pend_copy = pend_df.copy()
    if 'FECHA DE ATENCION' in pend_copy.columns:
        pend_copy['FECHA DE ATENCION'] = pend_copy['FECHA DE ATENCION'].dt.strftime('%Y-%m-%d')
    if 'HORA DE ATENCION' in pend_copy.columns:
        pend_copy['HORA DE ATENCION'] = pend_copy['HORA DE ATENCION'].astype(str)

    drive_json_data = {
        "origen": "BDAGENDAWEB + BDSVCES + BDSAP (SQL Server 172.200.6.135)",
        "total_historico_asignadas": summary_data.get("kpis", {}).get("total_asignadas", cache_pack.get("total_historico_asignadas", 0)),
        "total_historico_asistidas": summary_data.get("kpis", {}).get("asistidas", cache_pack.get("total_historico_asistidas", 0)),
        "total_historico_inasistidas": summary_data.get("kpis", {}).get("inasistidas", cache_pack.get("total_historico_inasistidas", 0)),
        "total_historias_pendientes": len(pend_copy),
        "sedes_summary": sedes_summary,
        "usuarios_activos": usuarios_df.fillna("").astype(str).to_dict(orient="records") if hasattr(usuarios_df, 'to_dict') and not usuarios_df.empty else [],
        "historias_pendientes": pend_copy.to_dict(orient="records") if not pend_copy.empty else []
    }

    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(drive_json_data, f, ensure_ascii=False, indent=2)

    json_size_mb = os.path.getsize(json_path) / (1024*1024)
    print(f"1. Caché JSON para Google Drive generado: {json_path} ({json_size_mb:.2f} MB)")


    # 2. Exportar a Excel Multi-Hoja para Google Drive / Google Sheets
    xlsx_path = "Caché_Historias_Pendientes_AGENDAWEB.xlsx"
    wb = Workbook()
    
    # Hoja 1: Resumen por Sedes
    ws1 = wb.active
    ws1.title = "Resumen por Sedes"
    ws1.append(["Sede / Médico", "Total Asignadas", "Asistidas", "Inasistidas", "Pendientes", "% Pendientes", "% Inasistidas"])
    for s in sedes_summary:
        s_name = s.get("NOMBRE IPS") or s.get("sede") or ""
        ws1.append([s_name, s.get("total"), s.get("asistidas"), s.get("inasistidas"), s.get("pendientes"), s.get("pct_pendientes"), s.get("pct_inasistidas")])
        for m in s.get("medicos", []):
            m_name = "   └─ " + (m.get("profesional") or m.get("NOMBRE PROFESIONAL") or "")
            ws1.append([m_name, m.get("total"), m.get("asistidas"), m.get("inasistidas"), m.get("pendientes"), m.get("pct_pendientes"), m.get("pct_inasistidas")])
            for f in m.get("fechas", []):
                f_name = "      └─ " + (f.get("periodo") or f"{f.get('ano')} - {f.get('mes')}")
                ws1.append([f_name, f.get("total"), f.get("asistidas"), f.get("inasistidas"), f.get("pendientes"), f.get("pct_pendientes"), f.get("pct_inasistidas")])



    # Hoja 2: Historias Pendientes (Asistidas SI, Atendidas NO)
    ws2 = wb.create_sheet(title="Historias Pendientes")
    ws2.append(["ID Cita", "Fecha", "Hora", "Sede", "Tipo Doc", "Documento", "Paciente", "Programa", "Profesional"])
    for _, r in pend_df.iterrows():
        f_str = r['FECHA DE ATENCION'].strftime('%d/%m/%Y') if hasattr(r['FECHA DE ATENCION'], 'strftime') else str(r['FECHA DE ATENCION'])
        ws2.append([
            str(r.get("ID_CITA")), f_str, str(r.get("HORA DE ATENCION")),
            r.get("NOMBRE IPS"), str(r.get("ID", "CC")), str(r.get("DOCUMENTO")),
            r.get("NOMBRE PACIENTE"), r.get("PROGRAMA"), r.get("NOMBRE PROFESIONAL")
        ])

    # Hoja 3: Usuarios Activos
    ws3 = wb.create_sheet(title="Usuarios Activos")
    ws3.append(["Profesional", "Cédula / Identificación", "Sede Asignada"])
    for _, r in usuarios_df.iterrows():
        ws3.append([r.get("NOMBRE PROFESIONAL"), str(r.get("Cedula")), r.get("NOMBRE IPS")])

    wb.save(xlsx_path)
    xlsx_size_mb = os.path.getsize(xlsx_path) / (1024*1024)
    print(f"2. Libro Excel para Google Drive generado: {xlsx_path} ({xlsx_size_mb:.2f} MB)")

    print(f"\n Archivos listos para sincronizar con Google Drive en Google Drive Desktop!")
